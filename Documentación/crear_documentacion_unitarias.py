from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name('Documentacion_Pruebas_Unitarias_ShipTranslate.docx')

BLUE = '2E74B5'
DARK_BLUE = '1F4D78'
INK = '243447'
MUTED = '667085'
LIGHT_BLUE = 'E8EEF5'
LIGHT_GRAY = 'F2F4F7'
PALE_GREEN = 'EAF5EE'
GREEN = '237A46'
PALE_GOLD = 'FFF5D6'
BORDER = 'B8C3CE'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color=BORDER, size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), size)
        node.set(qn('w:color'), color)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(total))
    tbl_w.set(qn('w:type'), 'dxa')

    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), str(indent))
    tbl_ind.set(qn('w:type'), 'dxa')

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(width))
            tc_w.set(qn('w:type'), 'dxa')
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def style_run(run, size=10, bold=False, color=INK, italic=False, font='Calibri'):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), font)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement('w:fldChar')
    fld_char.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_char, instr, separate, text, end])
    style_run(run, size=9, color=MUTED)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f'Heading {level}')
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        style_run(first, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        style_run(rest)
    else:
        style_run(p.add_run(text))
    return p


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run(f'{label}: '), size=10.5, bold=True, color=INK)
    style_run(p.add_run(value), size=10.5, color=INK)


def add_callout(doc, title, body, fill=LIGHT_BLUE, title_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    style_run(p.add_run(title), size=10.5, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    style_run(p2.add_run(body), size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths, font_size=8.5, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    set_table_geometry(table, widths)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(text), size=8.5, bold=True, color=DARK_BLUE)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                set_cell_shading(cell, 'F8FAFC')
        for index, value in enumerate(values):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, status_col) else WD_ALIGN_PARAGRAPH.LEFT
            color = GREEN if status_col is not None and index == status_col else INK
            bold = status_col is not None and index == status_col
            style_run(p.add_run(str(value)), size=font_size, bold=bold, color=color)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in [
    ('Heading 1', 16, BLUE, 16, 8),
    ('Heading 2', 13, BLUE, 12, 6),
    ('Heading 3', 12, DARK_BLUE, 8, 4),
]:
    style = styles[name]
    style.font.name = 'Calibri'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
style_run(hp.add_run('SHIPTRANSLATE  |  DOCUMENTACIÓN DE PRUEBAS'), size=8.5, bold=True, color=MUTED)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
style_run(fp.add_run('ShipTranslate  •  Página '), size=9, color=MUTED)
add_field(fp, 'PAGE')

# Masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(3)
style_run(p.add_run('DOCUMENTACIÓN TÉCNICA'), size=10, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
style_run(p.add_run('Pruebas unitarias realizadas'), size=25, bold=True, color=INK)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
style_run(p.add_run('Sistema ShipTranslate · Backend, parsing EDI/PDF y validación de tracking'), size=12.5, color=MUTED)

for label, value in [
    ('Proyecto', 'ShipTranslate / Proyecto Tranship'),
    ('Fecha de ejecución', '24 de junio de 2026'),
    ('Versión del documento', '1.0'),
    ('Estado', '18 de 18 pruebas aprobadas'),
]:
    add_label_value(doc, label, value)

doc.add_paragraph().paragraph_format.space_after = Pt(1)
add_callout(
    doc,
    'Resultado general: APROBADO',
    'La batería se ejecutó sin fallos: 13 pruebas de backend y 5 pruebas de frontend. El build de producción del frontend también finalizó correctamente.',
    fill=PALE_GREEN,
    title_color=GREEN,
)

add_heading(doc, '1. Objetivo', 1)
add_body(doc, 'Verificar de manera automática y repetible las reglas de negocio con mayor impacto en la interpretación de archivos EDI y BL: extracción de información estructurada, selección de ETA, normalización de fechas, conversión de toneladas métricas a kilogramos y validación de campos obligatorios en la vista de tracking.')
add_body(doc, 'Estas pruebas constituyen una línea base de regresión. Su propósito es detectar rápidamente cambios que alteren resultados ya validados antes de desplegar nuevas versiones del sistema.')

add_heading(doc, '2. Alcance y enfoque', 1)
add_body(doc, 'Tipo de prueba: pruebas unitarias automatizadas ejecutadas con Vitest. Cada caso comprueba una función o regla aislada mediante datos controlados y aserciones determinísticas.')
add_table(doc, ['Incluido', 'Fuera de esta ejecución'], [
    ('Parser de mensajes EDIFACT y selección de TDT/DTM.', 'OCR real sobre imágenes o PDF escaneado.'),
    ('Fechas de emisión de BL y conversión de pesos.', 'Consumo real de OpenAI o servicios externos.'),
    ('Validación obligatoria para BL tramp y escaneados.', 'Persistencia real contra MongoDB.'),
    ('Compilación de producción del frontend.', 'Pruebas de carga, seguridad o concurrencia.'),
], [4680, 4680], font_size=9)

add_heading(doc, '3. Ambiente y herramientas', 1)
add_table(doc, ['Componente', 'Versión / uso'], [
    ('Node.js', 'v24.15.0'),
    ('npm', '11.12.1'),
    ('Vitest', '4.1.9 · ejecución de pruebas unitarias'),
    ('@vitest/coverage-v8', '4.1.9 · medición de cobertura'),
    ('Supertest', '7.2.2 · instalado para futuras pruebas HTTP'),
    ('Testing Library React', '16.3.2 · disponible para pruebas de componentes'),
    ('jsdom', '29.1.1 · entorno DOM disponible para frontend'),
    ('Sistema operativo', 'Windows / PowerShell'),
], [3000, 6360], font_size=9)

doc.add_page_break()
add_heading(doc, '4. Resumen de ejecución', 1)
add_table(doc, ['Área', 'Archivos', 'Pruebas', 'Aprobadas', 'Fallidas', 'Estado'], [
    ('Backend', '2', '13', '13', '0', 'APROBADO'),
    ('Frontend', '1', '5', '5', '0', 'APROBADO'),
    ('Total', '3', '18', '18', '0', 'APROBADO'),
], [2400, 1100, 1300, 1500, 1300, 1760], font_size=9, status_col=5)

add_heading(doc, '5. Casos de prueba del backend', 1)
add_heading(doc, '5.1 Parser EDI', 2)
edi_rows = [
    ('UT-BE-EDI-001', 'Datos principales', 'EDI con UNB, BGM, RFF, NAD, LOC, GID, FTX y MEA.', 'Extrae intercambio, partes, puertos, carga y peso.', 'APROBADO'),
    ('UT-BE-EDI-002', 'ETA del tramo principal', 'Mensaje con TDT+20 y TDT+40, ambos con DTM+132.', 'Selecciona viaje, nave y ETA asociados al TDT+40.', 'APROBADO'),
    ('UT-BE-EDI-003', 'EDI sin mensajes', 'Intercambio sin segmento UNH.', 'Devuelve arreglo vacío y total de mensajes igual a cero.', 'APROBADO'),
]
add_table(doc, ['ID', 'Funcionalidad', 'Entrada / condición', 'Resultado esperado', 'Estado'], edi_rows, [1350, 1700, 2350, 2700, 1260], font_size=7.8, status_col=4)

add_heading(doc, '5.2 Fechas de emisión del BL', 2)
date_rows = [
    ('UT-BE-FEC-001', 'Apr 20, 2026', '20-04-2026', 'APROBADO'),
    ('UT-BE-FEC-002', 'September 5, 2025', '05-09-2025', 'APROBADO'),
    ('UT-BE-FEC-003', 'Texto con Jan 1, 2027', '01-01-2027', 'APROBADO'),
    ('UT-BE-FEC-004', 'Texto sin fecha compatible', 'Cadena vacía', 'APROBADO'),
]
add_table(doc, ['ID', 'Entrada', 'Resultado esperado', 'Estado'], date_rows, [1700, 3000, 3000, 1660], font_size=8.5, status_col=3)

add_heading(doc, '5.3 Conversión y priorización de pesos', 2)
weight_rows = [
    ('UT-BE-PES-001', '48,030 toneladas', '48030 kg', 'APROBADO'),
    ('UT-BE-PES-002', '532.110 toneladas', '532110 kg', 'APROBADO'),
    ('UT-BE-PES-003', '0.5 toneladas', '500 kg', 'APROBADO'),
    ('UT-BE-PES-004', 'Peso bruto y neto presentes', 'Prioriza TOTAL GROSS WEIGHT', 'APROBADO'),
    ('UT-BE-PES-005', 'MT, M/T y METRIC TONS', 'Convierte todas las variantes a KG', 'APROBADO'),
    ('UT-BE-PES-006', 'Objeto con arreglos anidados y M.T.', 'Normaliza todas las cadenas sin alterar números', 'APROBADO'),
]
add_table(doc, ['ID', 'Entrada / condición', 'Resultado esperado', 'Estado'], weight_rows, [1700, 3000, 3000, 1660], font_size=8.2, status_col=3)

doc.add_page_break()
add_heading(doc, '6. Casos de prueba del frontend', 1)
front_rows = [
    ('UT-FE-VAL-001', 'BL escaneado completo', 'Todos los campos obligatorios tienen contenido.', 'No informa campos pendientes.', 'APROBADO'),
    ('UT-FE-VAL-002', 'Campo compuesto por espacios', 'Shipper contiene únicamente espacios.', 'Lo considera vacío y obligatorio pendiente.', 'APROBADO'),
    ('UT-FE-VAL-003', 'Freight payable at vacío', 'Campo opcional sin contenido.', 'Permite continuar sin marcarlo como faltante.', 'APROBADO'),
    ('UT-FE-VAL-004', 'Segundo BL tramp incompleto', 'Date of issue vacío en uno de dos BL.', 'Detecta que existe al menos un BL incompleto.', 'APROBADO'),
    ('UT-FE-VAL-005', 'Opcionales tramp vacíos', 'Measurement y Prepaid at sin contenido.', 'No los considera obligatorios.', 'APROBADO'),
]
add_table(doc, ['ID', 'Funcionalidad', 'Entrada / condición', 'Resultado esperado', 'Estado'], front_rows, [1350, 1750, 2400, 2600, 1260], font_size=7.8, status_col=4)

add_heading(doc, '7. Cobertura obtenida', 1)
add_table(doc, ['Área medida', 'Sentencias', 'Ramas', 'Funciones', 'Líneas'], [
    ('Backend completo importado', '27,62 %', '20,35 %', '17,05 %', '27,90 %'),
    ('Utilidad frontend trackingValidation.js', '100 %', '75 %', '100 %', '100 %'),
], [3280, 1520, 1520, 1520, 1520], font_size=9)

add_callout(
    doc,
    'Interpretación de cobertura',
    'El porcentaje del backend es una línea base, no una cobertura total satisfactoria: al importar shipments.js también se contabilizan rutas, modelos, OCR, OpenAI y persistencia que aún no poseen pruebas aisladas. En frontend, el 100 % corresponde únicamente a la utilidad de validación creada, no a toda la aplicación React.',
    fill=PALE_GOLD,
    title_color='7A5A00',
)

add_heading(doc, '8. Hallazgo y corrección realizada', 1)
add_body(doc, 'Hallazgo: durante la primera ejecución falló UT-BE-PES-006. La expresión regular convertía MT, M/T y METRIC TONS, pero no reconocía correctamente M.T. cuando el punto final era seguido por un espacio o el fin de la cadena.')
add_body(doc, 'Causa técnica: se utilizaba un límite de palabra (\\b) después de una unidad terminada en punto. Este límite no se cumple entre dos caracteres no alfanuméricos.')
add_body(doc, 'Corrección: se reemplazó el límite de palabra por una condición negativa (?![A-Z]) en las funciones de búsqueda y conversión. Tras el ajuste se repitieron todas las pruebas y se obtuvieron 18 aprobaciones y 0 fallos.')

add_heading(doc, '9. Evidencia y trazabilidad', 1)
add_table(doc, ['Evidencia', 'Ubicación / resultado'], [
    ('Pruebas del parser EDI', 'BACKEND/tests/ediParser.test.js'),
    ('Pruebas de fecha y peso', 'BACKEND/tests/pdfHelpers.test.js'),
    ('Pruebas de validación', 'FRONTEND/src/utils/trackingValidation.test.js'),
    ('Utilidad validada', 'FRONTEND/src/utils/trackingValidation.js'),
    ('Ejecución backend', '2 archivos aprobados · 13 pruebas aprobadas'),
    ('Ejecución frontend', '1 archivo aprobado · 5 pruebas aprobadas'),
    ('Compilación frontend', 'vite build completado · 100 módulos transformados'),
], [3000, 6360], font_size=9)

add_heading(doc, '10. Reproducción de las pruebas', 1)
add_body(doc, 'Ejecutar desde PowerShell en el directorio raíz del proyecto:')
commands = [
    'cd BACKEND',
    'npm test',
    'npm run test:coverage',
    'cd ..\\FRONTEND',
    'npm test',
    'npm run test:coverage',
    'npm run build',
]
table = doc.add_table(rows=1, cols=1)
set_table_geometry(table, [9360])
set_table_borders(table, color='D0D5DD')
set_cell_shading(table.cell(0, 0), LIGHT_GRAY)
p = table.cell(0, 0).paragraphs[0]
p.paragraph_format.space_after = Pt(0)
for index, command in enumerate(commands):
    if index:
        p.add_run('\n')
    style_run(p.add_run(command), size=9, color=INK, font='Consolas')

add_heading(doc, '11. Criterios de aceptación', 1)
acceptance_rows = [
    ('Todas las pruebas prioritarias deben finalizar sin fallos.', 'Cumplido'),
    ('La ETA debe provenir del TDT de etapa mayor.', 'Cumplido'),
    ('Las variantes de toneladas métricas deben convertirse a KG.', 'Cumplido'),
    ('Los campos opcionales no deben bloquear el guardado.', 'Cumplido'),
    ('El frontend debe compilar después de la modularización.', 'Cumplido'),
]
add_table(doc, ['Criterio', 'Resultado'], acceptance_rows, [7300, 2060], font_size=9, status_col=1)

add_heading(doc, '12. Riesgos pendientes y próximos pasos', 1)
add_body(doc, 'La batería actual protege reglas puras, pero todavía no cubre completamente los flujos que dependen de infraestructura. La siguiente iteración debería incorporar pruebas HTTP con Supertest, base de datos temporal o mocks de Mongoose, autenticación y permisos por rol, persistencia de trackings, solicitudes administrativas y pruebas de componentes React para alertas y redirecciones.')
add_body(doc, 'También se recomienda crear fixtures anonimizados de BL tramp y escaneados. Esos archivos permitirían probar regresiones del parser sin depender de documentos productivos ni realizar llamadas reales a OpenAI.')

add_heading(doc, '13. Conclusión', 1)
add_body(doc, 'La ejecución demuestra que las reglas unitarias seleccionadas funcionan de acuerdo con los resultados esperados. Las 18 pruebas aprobadas proporcionan una base repetible para continuar desarrollando el sistema con menor riesgo de regresión. Además, la detección del formato M.T. confirma el valor práctico de automatizar casos derivados de documentos reales.')

# Document properties
doc.core_properties.title = 'Documentación de pruebas unitarias - ShipTranslate'
doc.core_properties.subject = 'Evidencia de ejecución de pruebas unitarias'
doc.core_properties.author = 'Proyecto ShipTranslate'
doc.core_properties.keywords = 'pruebas unitarias, Vitest, EDI, PDF, BL, tracking'

doc.save(OUT)
print(OUT)
